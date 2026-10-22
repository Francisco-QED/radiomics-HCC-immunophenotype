"""Results deck v2: follows the structure and table styles of updates-08-17-26.pptx (title, Patients, methods, cohort characteristics, then one section per endpoint:
data split table, performance tables, ROC / KM + time-dependent AUC, SHAP). Starts from that deck (theme kept, first slide kept, the rest removed); never overwrites it.

Performance tables carry only the main AUC (classification) / C-index (recurrence) with 95% CIs, plus lock sensitivity / specificity for the binary endpoints.
No OOF sensitivity / specificity and no time-dependent AUC tables (the time-dependent AUC appears only in the KM figures).
Sections are added only when their inputs exist.
"""
import argparse, glob, sys
from datetime import date
from pathlib import Path
sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))
import numpy as np, pandas as pd
from scipy.stats import chi2_contingency, fisher_exact, mannwhitneyu
from sklearn.metrics import roc_curve
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
import build_deck as bd
import build_presentation_tables as bt

ROOT = bd.ROOT
OUT = ROOT / "Results/main_pipeline_claude/presentation_20261020"
FAM = {"ridge": "Ridge", "lasso": "Lasso", "elastic_net": "Elastic net"}
CAT = [("radiomics", "Radiomics", "#1f77b4"), ("radiomics_clinical", "Radiomics + clinico-radiologic", "#d62728"), ("clinical", "Clinico-radiologic", "#2ca02c")]
ENDPOINTS = {
    "tils_ge2": dict(label="TILs 2-3 vs 0-1", short="TILs", dir=OUT / "tils_ge2", prefix="tils_ge2",
                     oof={"radiomics": ROOT / "Results/main_pipeline_claude/main_pipeline_tils/ladder_post_20261021/tilsGE2_post_*/oof_predictions.csv", "radiomics_clinical": ROOT / "Results/main_pipeline_claude/main_pipeline_tils/ladder_clinrad_post_20261021/tilsGE2cr_post_*/oof_predictions.csv",
                          "clinical": ROOT / "Results/main_pipeline_claude/main_pipeline_tils/clinical_only_v4_20261021/oof_predictions.csv"}),
    "tils_ge3": dict(label="TILs 3 vs 0-2", short="TILs3", dir=OUT / "tils_ge3", prefix="tils_ge3",
                     oof={"radiomics": ROOT / "Results/main_pipeline_claude/main_pipeline_tils/ge3_ladder_post_20261021/tilsGE3_post_*/oof_predictions.csv", "radiomics_clinical": ROOT / "Results/main_pipeline_claude/main_pipeline_tils/ge3_ladder_clinrad_post_20261021/tilsGE3cr_post_*/oof_predictions.csv",
                          "clinical": ROOT / "Results/main_pipeline_claude/main_pipeline_tils/ge3_clinical_only_post_20261021/oof_predictions.csv"}),
    "tls_ge2": dict(label="TLS 2 vs 0-1", short="TLS", dir=OUT / "tls_ge2", prefix="tls_ge2",
                    oof={"radiomics": ROOT / "Results/main_pipeline_claude/main_pipeline_tls/ladder_post_20261021/tlsGE2_post_*/oof_predictions.csv", "radiomics_clinical": ROOT / "Results/main_pipeline_claude/main_pipeline_tls/ladder_clinrad_post_20261021/tlsGE2cr_post_*/oof_predictions.csv",
                         "clinical": ROOT / "Results/main_pipeline_claude/main_pipeline_tls/clinical_only_post_20261021/oof_predictions.csv"}),
}


def ci(v, lo, hi, p=2):
    return f"{v:.{p}f} ({lo:.{p}f}-{hi:.{p}f})"


# ---------------------------------------------------------------- cohort characteristics (whole cohort, compact, slide-17 style)
def cat_str(s, order=None):
    d = s.dropna(); n = len(d)
    if n == 0: return "-"
    vc = d.value_counts(); levels = [l for l in (order or []) if l in vc.index] + [l for l in vc.index if l not in (order or [])]
    return " / ".join(f"{l} {vc[l]} ({100 * vc[l] / n:.1f}%)" for l in levels), n


def cont_str(x, med=False):
    x = x.dropna()
    return (f"{x.median():.1f} [IQR {x.quantile(.25):.1f}-{x.quantile(.75):.1f}]" if med else f"{x.mean():.1f} ± {x.std():.1f}"), len(x)


def cohort_slide(prs, D):
    def rows(spec):
        out = []
        for name, col, kind in spec:
            if kind == "cat": txt, n = cat_str(D[col], bt.ORDER.get(col))
            else: txt, n = cont_str(D[col], med=(kind == "med"))
            out.append({"Variable": f"{name} (n={n})", "Value": txt})
        return pd.DataFrame(out)
    mri = rows([("Morphology", "Morphology", "cat"), ("Capsule (on PVP)", "Capsule (on PVP)", "cat"), ("Margin", "Margin", "cat"), ("Mosaic appearance", "Mosaic appearance", "cat"), ("Internal arteries", "Internal arteries", "cat"),
                ("Macrovascular invasion", "Macrovascular invasion", "cat"), ("Fat content", "Fat content", "cat"), ("T2 signal", "T2 signal", "cat"), ("Restricted diffusion", "Restricted diffusion", "cat"),
                ("Enhancement (arterial phase)", "Arterial enhancement", "cat"), ("LI-RADS v2018", "LI-RADS v2018", "cat"), ("Maximum tumor size on T1 PVP, cm", "Maximum tumor size on T1 PVP, cm", "mean"), ("HBP signal", "HBP signal", "cat")])
    demo = rows([("Age, years", "Age, years", "mean"), ("Gender", "Gender", "cat"), ("BMI, kg/m²", "BMI, kg/m2", "mean"), ("Race", "Race", "cat"), ("Ethnicity", "Ethnicity", "cat")])
    clin = rows([("Etiology", "Etiology of liver disease", "cat"), ("Cirrhosis", "Cirrhosis", "cat"), ("Serum AFP, ng/mL", "Serum AFP, ng/mL", "med"), ("Treatment", "Treatment cohort", "cat")])
    s = bd.new_slide(prs, "Clinical/demographic + MRI characteristics")
    bd.text_box(s, 0.5, 0.95, 5.8, 0.3, [("Qualitative MRI (SA2 read on post-treatment MRI)",)], size=11)
    bd.add_table(s, mri, 0.5, 1.3, [2.2, 3.9], font=7.5, row_h=0.3, align_center_from=9)
    bd.text_box(s, 6.67, 0.95, 6.4, 0.3, [("Demographics",)], size=11)
    bd.add_table(s, demo, 6.67, 1.3, [1.7, 4.7], font=7.5, row_h=0.3, align_center_from=9)
    bd.text_box(s, 6.67, 3.55, 6.4, 0.3, [("Clinical",)], size=11)
    bd.add_table(s, clin, 6.67, 3.9, [1.7, 4.7], font=7.5, row_h=0.3, align_center_from=9)
    bd.footnote(s, "Entire cohort, N=141 (SA1 untreated n=111; SA2 immunotherapy n=30). Categorical: n (%) of non-missing; continuous: mean ± SD (AFP: median [IQR]). Variable sources: mastersheet tabs SA1 + SA2-post.", top=6.9)


# ---------------------------------------------------------------- dev / lock split table (slide-18 style)
def split_table(D, split, spec, n_dev, n_lock):
    dev_ids, lock_ids = split.index[split["set"] == "dev"], split.index[split["set"] == "lock"]; rows = []
    for name, col, kind in spec:
        a, b = D.loc[dev_ids, col], D.loc[lock_ids, col]
        if kind == "cont":
            f = lambda x: (lambda x: f"{x.median():.1f} [{x.quantile(.25):.1f} - {x.quantile(.75):.1f}]")(x.dropna())
            p = mannwhitneyu(a.dropna(), b.dropna()).pvalue; test = "Mann-Whitney U"; A, B = f(a), f(b)
        else:
            levels = sorted(set(a.dropna()) | set(b.dropna()), key=lambda l: -((a == l).sum() + (b == l).sum()))
            fmt = lambda x: ", ".join(f"{l}: {(x == l).sum()} ({100 * (x == l).sum() / x.notna().sum():.1f}%)" for l in levels)
            T = np.array([[(a == l).sum() for l in levels], [(b == l).sum() for l in levels]])
            if len(levels) == 2: p = fisher_exact(T.T)[1]; test = "Fisher's exact"
            else: p = chi2_contingency(T, correction=False)[1]; test = "Chi-Square"
            A, B = fmt(a), fmt(b)
        rows.append({"Variable": name, f"Development set (N={n_dev})": A, f"Lock set (N={n_lock})": B, "p-value": "<0.001" if p < 0.001 else f"{p:.3f}", "Statistical test": test})
    return pd.DataFrame(rows)


def split_slide(prs, title, df, note=""):
    s = bd.new_slide(prs, title); bd.add_table(s, df, 0.45, 1.1, [2.9, 3.3, 3.3, 0.8, 1.5], font=9, row_h=0.45, align_center_from=1, header_h=0.5)
    bd.footnote(s, "Development vs lock set; the lock set is held out from all training, tuning and model selection. " + note, top=6.9)


# ---------------------------------------------------------------- performance tables
def best_per_combo(df, key):
    return df.loc[df.groupby("sequence_set")[key].idxmax()].sort_values(key, ascending=False).reset_index(drop=True)


def cls_table(df, n=None):
    rows = []
    for i, r in df.head(n).iterrows() if n else df.iterrows():
        deg = min(r.lock_sens, r.lock_spec) <= 0.05
        rows.append({"Sequences": r.sequence_set.replace("+", " + ") if r.sequence_set != "clinicoradiologic" else "all variables", "Model": FAM[r.family] + (" †" if deg else ""),
                     "OOF AUC (95% CI)": ci(r.oof_auc_avgprob, r.oof_auc_ci_lo, r.oof_auc_ci_hi), "Lock AUC (95% CI)": ci(r.lock_auc, r.lock_auc_ci_lo, r.lock_auc_ci_hi),
                     "Lock sensitivity (95% CI)": ci(r.lock_sens, r.lock_sens_ci_lo, r.lock_sens_ci_hi), "Lock specificity (95% CI)": ci(r.lock_spec, r.lock_spec_ci_lo, r.lock_spec_ci_hi)})
    return pd.DataFrame(rows)


def rec_table(df, n=None):
    rows = []
    for i, r in df.head(n).iterrows() if n else df.iterrows():
        rows.append({"Sequences": r.sequence_set.replace("+", " + ") if r.sequence_set != "clinical" else "all variables", "Model": FAM[r.family],
                     "OOF C-index (95% CI)": ci(r.oof_c_avg, r.oof_c_ci_lo, r.oof_c_ci_hi), "Lock C-index (95% CI)": ci(r.lock_c, r.lock_c_ci_lo, r.lock_c_ci_hi)})
    return pd.DataFrame(rows)


def perf_slide(prs, title, df, kind, foot):
    s = bd.new_slide(prs, title); n = len(df)
    widths = [2.3, 1.3, 2.2, 2.2, 2.3, 2.3] if kind == "cls" else [2.8, 1.8, 3.4, 3.4]
    bd.add_table(s, df, 0.45 if kind == "cls" else 0.9, 1.1, widths, font=10 if n <= 10 else 9, row_h=0.4 if n <= 10 else 0.34, align_center_from=2, header_h=0.5)
    bd.footnote(s, foot, top=6.85)


CLS_FOOT = ("One model per sequence combination: the family (ridge / lasso / elastic net) with the highest mean OOF AUC over 10 repeated 5-fold splits (development data only). OOF AUC = AUC of the split-averaged out-of-fold predictions, bootstrap 95% CI. "
            "Lock sensitivity / specificity at the Youden threshold of the development OOF predictions (Wilson CI). † = (almost) everyone classified one way at that threshold. Lock set n=30.")
REC_FOOT = ("One model per sequence combination: the family with the highest mean OOF C-index over 10 repeated 5-fold splits (development data only). OOF C-index = Harrell C of the split-averaged out-of-fold risk, patient-bootstrap 95% CI. "
            "Lock C-index from the final development-fit model (n=30). The lock set was opened earlier in this project: lock results are exploratory.")


def roc_figure(ep, best, path):
    fig, ax = plt.subplots(1, 2, figsize=(12.5, 4.9)); handles = []
    for tag, name, col in CAT:
        r = best[tag]; o = pd.concat([pd.read_csv(f) for f in glob.glob(str(ep["oof"][tag]))]); o = o[(o.family == r.family) & o.prob.notna()]
        if tag != "clinical": o = o[o.sequence_set == r.sequence_set]
        avg = o.groupby("PatientID").agg(y=("y", "first"), prob=("prob", "mean")); fpr, tpr, _ = roc_curve(avg.y, avg.prob)
        label = f"{name} ({r.sequence_set.replace('+', ' + ') if tag != 'clinical' else 'all variables'}, {FAM[r.family].lower()})"; h, = ax[0].plot(fpr, tpr, color=col, lw=2, label=label); handles.append(h)
        folder = f"{r.sequence_set.replace('+', '_')}__{r.family}" if tag != "clinical" else f"clinicoradiologic__{r.family}"
        lp = pd.read_csv(ep["dir"] / "models" / tag / folder / "lock_predictions.csv"); fpr, tpr, _ = roc_curve(lp.y_lock, lp.prob); ax[1].plot(fpr, tpr, color=col, lw=2)
    for k, t in enumerate(("Development set: out-of-fold (averaged over splits)", "Lock set (one-shot)")):
        ax[k].plot([0, 1], [0, 1], "--", color="gray", lw=1); ax[k].set_xlabel("1 - specificity"); ax[k].set_ylabel("Sensitivity"); ax[k].set_title(t, fontsize=11); ax[k].set_aspect("equal"); ax[k].spines[["top", "right"]].set_visible(False)
    fig.legend(handles=handles, loc="lower center", ncol=3, frameon=False, fontsize=9.5); plt.tight_layout(rect=(0, 0.07, 1, 1)); plt.savefig(path, dpi=200, bbox_inches="tight"); plt.close()


def classification_section(prs, key, D, split, spec_extra):
    ep = ENDPOINTS[key]; d = ep["dir"]; f = {t: d / f"{ep['prefix']}_model_table_{t}.csv" for t, _, _ in CAT}
    if not all(p.exists() for p in f.values()): print("skip", key, "(model tables missing)"); return None
    dev_n, lock_n = int((split["set"] == "dev").sum()), int((split["set"] == "lock").sum())
    spec = [("Age", "Age, years", "cont"), ("Gender", "Gender", "cat"), ("BMI", "BMI, kg/m2", "cont"), ("Cirrhosis", "Cirrhosis", "cat"), ("Capsule (on PVP yes/no)", "Capsule (on PVP)", "cat"),
            ("Mosaic appearance (yes/no)", "Mosaic appearance", "cat"), ("Macrovascular invasion (yes/no)", "Macrovascular invasion", "cat"), ("Maximum tumor size on T1 portal venous phase (PVP)", "Maximum tumor size on T1 PVP, cm", "cont")] + spec_extra
    split_slide(prs, f"{ep['label']} - development/testing data split", split_table(D, split, spec, dev_n, lock_n))
    rad, com, cli = (pd.read_csv(f[t]) for t, _, _ in CAT)
    rb, cb = best_per_combo(rad, "oof_auc_repeat_mean"), best_per_combo(com, "oof_auc_repeat_mean"); clb = cli.sort_values("oof_auc_repeat_mean", ascending=False).reset_index(drop=True)
    perf_slide(prs, f"{ep['label']} - radiomics models, top 10 of 14 sequence combinations", cls_table(rb, 10), "cls", CLS_FOOT)
    perf_slide(prs, f"{ep['label']} - radiomics + clinico-radiologic variables, all 14 combinations", cls_table(cb), "cls", CLS_FOOT)
    best = {"radiomics": rb.iloc[0], "radiomics_clinical": cb.iloc[0], "clinical": clb.iloc[0]}
    s = bd.new_slide(prs, f"{ep['label']} - best model of each category")
    t = cls_table(pd.DataFrame([best[t] for t, _, _ in CAT])); t.insert(0, "Category", [n for _, n, _ in CAT])
    bd.add_table(s, t, 0.35, 1.05, [2.4, 1.4, 1.0, 1.9, 1.9, 2.0, 2.0], font=9, row_h=0.4, align_center_from=3, header_h=0.5)
    roc_figure(ep, best, d / "roc_best_per_category_v2.png"); pic = bd.picture(s, d / "roc_best_per_category_v2.png", 0.9, 2.95, height=4.0); pic.left = int((prs.slide_width - pic.width) / 2)
    bd.footnote(s, "Best model per category selected on development OOF AUC. Left: OOF predictions averaged over the 10 splits (n=%d). Right: one-shot lock set (n=%d). †: see previous slides." % (dev_n, lock_n), top=7.0)
    s = bd.new_slide(prs, f"{ep['label']} - SHAP explainability, best model of each category"); x = 0.35
    for tag, name, _ in CAT:
        r = best[tag]; folder = f"{r.sequence_set.replace('+', '_')}__{r.family}" if tag != "clinical" else f"clinicoradiologic__{r.family}"; img = d / "models" / tag / folder / "shap_beeswarm.png"
        if img.exists(): bd.picture(s, img, x, 1.15, width=4.2)
        x += 4.3
    bd.footnote(s, "Linear SHAP values on the final development-fit models (standardized selected features, development patients). Plots for all models: " + str(d / "models"), top=6.95)
    return best


# ---------------------------------------------------------------- recurrence
def recurrence_section(prs, D, split, r_best_dir):
    d = OUT / "recurrence"; f = {t: d / f"recurrence_model_table_{t}.csv" for t, _, _ in CAT}
    if not all(p.exists() for p in f.values()) or not (d / "km/km_summary.csv").exists(): print("skip recurrence"); return None
    dev_n, lock_n = int((split["set"] == "dev").sum()), int((split["set"] == "lock").sum())
    spec = [("Age", "Age, years", "cont"), ("Gender", "Gender", "cat"), ("Cirrhosis", "Cirrhosis", "cat"), ("Capsule (on PVP yes/no)", "Capsule (on PVP)", "cat"), ("Mosaic appearance (yes/no)", "Mosaic appearance", "cat"),
            ("Macrovascular invasion (yes/no)", "Macrovascular invasion", "cat"), ("Maximum tumor size on T1 portal venous phase (PVP)", "Maximum tumor size on T1 PVP, cm", "cont"), ("Serum AFP", "Serum AFP, ng/mL", "cont"),
            ("Recurrence", "Recurrence", "cat"), ("Follow-up time (days)", "Follow-up / time to recurrence, days", "cont")]
    split_slide(prs, "Recurrence - development/testing data split", split_table(D, split, spec, dev_n, lock_n), note="Follow-up time = time to recurrence or last follow-up. One patient without follow-up excluded.")
    rad, com, cli = (pd.read_csv(f[t]) for t, _, _ in CAT)
    rb, cb = best_per_combo(rad, "oof_c_repeat_mean"), best_per_combo(com, "oof_c_repeat_mean"); clb = cli.sort_values("oof_c_repeat_mean", ascending=False).reset_index(drop=True)
    perf_slide(prs, "Recurrence - radiomics models, top 10 of 14 sequence combinations", rec_table(rb, 10), "rec", REC_FOOT)
    perf_slide(prs, "Recurrence - radiomics + clinico-radiologic variables, all 14 combinations", rec_table(cb), "rec", REC_FOOT)
    best = {"radiomics": rb.iloc[0], "radiomics_clinical": cb.iloc[0], "clinical": clb.iloc[0]}; km = pd.read_csv(d / "km/km_summary.csv")
    for setname, title in (("lock", "lock set (exploratory)"), ("dev", "development set, out-of-fold risk")):
        s = bd.new_slide(prs, f"Recurrence - best model of each category: {title}"); x = 0.3
        for tag, name, _ in CAT:
            r = best[tag]; k = km[(km.tag == tag) & (km.set == setname)].iloc[0]
            row = pd.DataFrame([{"Model": f"{name}: {r.sequence_set.replace('+', ' + ') if tag != 'clinical' else 'all variables'}, {FAM[r.family].lower()}",
                                 "C-index (95% CI)": ci(r.oof_c_avg, r.oof_c_ci_lo, r.oof_c_ci_hi) if setname == "dev" else ci(r.lock_c, r.lock_c_ci_lo, r.lock_c_ci_hi), "Log-rank p": "<0.001" if k.logrank_p < 0.001 else f"{k.logrank_p:.3f}"}])
            bd.add_table(s, row, x, 1.05, [2.0, 1.2, 0.8], font=8, row_h=0.62, align_center_from=1, header_h=0.3)
            bd.picture(s, d / "km" / f"km_{tag}_{setname}.png", x, 2.05, width=4.0); bd.picture(s, d / "km" / f"tdauc_{tag}_{setname}.png", x + 0.25, 5.1, width=3.5); x += 4.35
        bd.footnote(s, "High / low risk = above / at-or-below the median risk of the set shown. Time-dependent AUC(t): censoring-weighted. " + ("Lock set was opened earlier in this project: exploratory." if setname == "lock" else "Models selected on this same development OOF risk: optimistic."), top=7.1)
    s = bd.new_slide(prs, "Recurrence - SHAP explainability, best model of each category"); x = 0.35
    for tag, name, _ in CAT:
        r = best[tag]; folder = f"{r.sequence_set.replace('+', '_')}__{r.family}" if tag != "clinical" else f"clinicoradiologic__{r.family}"; img = d / "models" / tag / folder / "shap_beeswarm.png"
        if img.exists(): bd.picture(s, img, x, 1.15, width=4.2)
        x += 4.3
    bd.footnote(s, "Linear SHAP values of the Cox linear predictor in the final development-fit models. Plots for all models: " + str(d / "models"), top=6.95)
    return best, km


def conclusions_slides(prs, out):
    def delta(a, b, k):
        d = b[k] - a[k]; return "similar to" if abs(d) < 0.02 else ("higher than" if d > 0 else "lower than")
    lines = []
    for key, nm in (("tils", "TILs 2-3 vs 0-1"), ("tils3", "TILs 3 vs 0-2"), ("tls", "TLS 2 vs 0-1")):
        b = out.get(key)
        if not b: continue
        r, c, k = b["radiomics"], b["radiomics_clinical"], b["clinical"]
        lines += [(nm,), f"Best radiomics model ({r.sequence_set.replace('+', ' + ')}, {FAM[r.family].lower()}): OOF AUC {ci(r.oof_auc_avgprob, r.oof_auc_ci_lo, r.oof_auc_ci_hi)}; lock AUC {ci(r.lock_auc, r.lock_auc_ci_lo, r.lock_auc_ci_hi)}.",
                  f"Adding clinico-radiologic variables: OOF AUC {delta(r, c, 'oof_auc_repeat_mean')} radiomics alone ({c.oof_auc_repeat_mean:.2f} vs {r.oof_auc_repeat_mean:.2f}). Clinico-radiologic alone: OOF AUC {k.oof_auc_repeat_mean:.2f}."]
    if out.get("rec"):
        b, _ = out["rec"]; r, c, k = b["radiomics"], b["radiomics_clinical"], b["clinical"]
        lines += [("Recurrence",), f"OOF C-index: radiomics {r.oof_c_repeat_mean:.2f} ({r.sequence_set.replace('+', ' + ')}), radiomics + clinico-radiologic {c.oof_c_repeat_mean:.2f}, clinico-radiologic alone {k.oof_c_repeat_mean:.2f}; all with overlapping 95% CIs.",
                  "Lock-set results for recurrence are exploratory (the lock set was opened earlier)."]
    lines += [("Caveats",), "Differences between model categories are small relative to the confidence intervals; the lock sets are small (30 patients; 7 TILs<2 and 20 TILs<3 negatives), so lock AUCs are imprecise.",
              "Estimates are averages over 10 repeated splits; single-split values (e.g. AUC ~0.69 for TP) fall in the upper tail of that distribution. One model per combination was chosen on development OOF performance."]
    s = bd.new_slide(prs, "Conclusions"); bd.text_box(s, 0.5, 1.0, 12.3, 5.9, lines, size=13, bullet=True)
    s = bd.new_slide(prs, "Next steps...")
    bd.text_box(s, 0.5, 1.0, 12.3, 5.0, ["Confirm the SA2 post-treatment qualitative MRI reads and extend the cohort (extended analysis, N=220) for external-style validation of the frozen models.", "Target IHC-derived endpoints.",
                                         "Target assessment of response to ICIs in SA2.", "Compare pre- vs post-treatment imaging for SA2 now that radiologic variables are post-treatment."], size=14, bullet=True)


def main():
    ap = argparse.ArgumentParser(); a = ap.parse_args()
    prs = Presentation(bd.TEMPLATE)
    ids = prs.slides._sldIdLst
    for sid in list(ids)[1:]:                       # keep the template's title slide
        prs.part.drop_rel(sid.rId); ids.remove(sid)
    s0 = prs.slides[0]
    bd.text_box(s0, 0.26, 5.6, 10, 0.5, [f"Results update, {date.today():%m/%d/%y}"], size=16)
    D, r = bt.build_frame(); tils = pd.read_csv(ROOT / "Results/main_pipeline_claude/main_pipeline/final/cohort_v3_141/split/split.csv").set_index("PatientID"); rec = pd.read_csv(ROOT / "Results/main_pipeline_claude/main_pipeline_recurrence/split/recurrence_split.csv").set_index("PatientID")
    s = bd.new_slide(prs, "Patients")
    pt = pd.DataFrame([{"Specific aim": "SA1", "Cohort size": "N = 111", "Description": "Retrospective cohort of resected, untreated HCC patients"},
                       {"Specific aim": "SA2", "Cohort size": "N = 30", "Description": "HCC patients who received neoadjuvant immunotherapy (IO, n=16; IO + SBRT, n=14) prior to resection"}])
    bd.add_table(s, pt, 0.5, 1.0, [2.0, 2.0, 7.6], font=12, row_h=0.5, align_center_from=1)
    bd.text_box(s, 0.5, 2.5, 11.5, 3.5, [("Inclusion criteria",), "Adult patients with chronic liver disease of any etiology and resected HCC (from last 14 years)", "Contrast-enhanced MRI (using gadoxetate) within 3 months of resection", "HCC lesions ≥ 2 cm",
                                         "Treatment with neoadjuvant IO (SA2 only)", ("Exclusion criteria",), "Poor MR image quality; previous locoregional or systemic therapy; unavailable tissue; previous liver transplant"], size=13, bullet=True)
    s = bd.new_slide(prs, "Data split and training")
    bd.text_box(s, 0.5, 1.1, 12.3, 5.6, [("Split",), "TILs / TLS: 111 / 30 patients in development / lock sets (stratified on outcome x cohort; same split for both endpoints). Recurrence: 110 / 30 (stratified on event x cohort; one patient without follow-up excluded).",
                                         "The lock set is never used for training, tuning or model selection and is scored once.", ("Validation (development set)",), "10 x repeated stratified 5-fold cross-validation (inner 3-fold for tuning); every preprocessing, selection and tuning step is refit inside each training fold.",
                                         "Performance = pooled out-of-fold (OOF) predictions: AUC (classification) or Harrell C-index (recurrence), with 95% CIs.", ("Endpoints",), "TILs 2-3 vs 0-1;  TILs 3 vs 0-2;  TLS 2 vs 0-1;  recurrence (time to recurrence; 2-year recurrence for ROC curves)."], size=14, bullet=True)
    s = bd.new_slide(prs, "Radiomics model development")
    bd.text_box(s, 0.5, 1.0, 7.9, 5.9, [("Per training fold",), "963 original PyRadiomics features (9 sequences x 107), per-sequence bin widths, N4-corrected images; SA2 imaging is post-treatment.", "Inter-reader filter: ICC(3,1) ≥ 0.85 on 23 patients segmented by two readers (estimated once on development data).",
                                         "QC: ≥ 80% observed, CV > 0.10; median imputation.", "Within-sequence complete-linkage clustering, |Spearman ρ| ≥ 0.85; one outcome-ranked representative per cluster.",
                                         "Ridge (top 10 / 15 / 20 clusters), lasso and elastic net (10-20 non-zero coefficients) logistic / Cox models; regularization tuned by inner 3-fold CV.",
                                         ("Clinico-radiologic variables",), "Age, sex, BMI, etiology, cirrhosis, log AFP, tumor size, qualitative MRI (capsule, margin, mosaic, internal arteries, fat, T2, HBP, LI-RADS), treatment (untreated / IO / IO + SBRT). Binary variables with minority < 10% dropped.",
                                         "Combined models: clinical variables added to the selected radiomics features (ridge) or to the candidate pool (lasso / elastic net)."], size=12, bullet=True)
    combos = pd.DataFrame({"Sequence combinations (14)": ["TP", "PVP", "DWI (b2)", "HBP", "TP + PVP", "TP + DWI", "TP + HBP", "PVP + DWI", "PVP + HBP", "DWI + HBP", "TP + PVP + DWI", "TP + PVP + HBP", "TP + DWI + HBP", "PVP + DWI + HBP"]})
    bd.add_table(s, combos, 9.0, 1.1, [3.5], font=10, row_h=0.33, align_center_from=9)
    cohort_slide(prs, D)
    out = {}
    out["tils"] = classification_section(prs, "tils_ge2", D, tils, [("TILs_resection", "TILs score", "cat"), ("TLS_resection", "TLS score", "cat")])
    out["tils3"] = classification_section(prs, "tils_ge3", D, tils, [("TILs_resection", "TILs score", "cat"), ("TLS_resection", "TLS score", "cat")])
    out["tls"] = classification_section(prs, "tls_ge2", D, tils, [("TLS_resection", "TLS score", "cat"), ("TILs_resection", "TILs score", "cat")])
    out["rec"] = recurrence_section(prs, D, rec, None)
    conclusions_slides(prs, out)
    dest = OUT / f"updates-{date.today():%m-%d-%y}_v2.pptx"; prs.save(dest); print("saved", dest, "slides:", len(prs.slides))
    return out


if __name__ == "__main__":
    main()

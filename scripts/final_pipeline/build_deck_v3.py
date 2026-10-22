"""Results deck v3: same analysis as v2 with (i) tumor grade (mastersheet column BM) added to the clinico-radiologic variables,
(ii) 'best model' = mean of OOF performance and lock performance (AUC; C-index for recurrence), (iii) mean-|SHAP| feature tables + complete feature lists,
(iv) a model-selection slide and (v) a with/without-grade sensitivity slide. Starts from the 08-17-26 deck (theme + title slide kept); never overwrites it.

Selection score = (OOF AUC + lock AUC) / 2, where OOF AUC is the AUC of the split-averaged out-of-fold predictions (the value shown in the tables); for recurrence the C-index analogue.
The score is used (a) to pick the family for each sequence combination, (b) to rank combinations, (c) to pick the best model of each category.
Because the lock set enters the score, lock numbers of selected models are selected estimates, not independent validation (stated on the slides).
Radiomics-only results are unchanged from the earlier run (grade does not enter them) and are read from presentation_20261020.
"""
import argparse, glob, json, re, subprocess, sys
from datetime import date
from pathlib import Path
sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))
import numpy as np, pandas as pd
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from sklearn.metrics import roc_curve
from pptx import Presentation
import build_deck as bd
import build_presentation_tables as bt
import build_deck_v2 as v2
import run_guard

BASE = bd.ROOT / "Results/main_pipeline_claude"
OLD, NEW = BASE / "presentation_20261020", BASE / "presentation_20261022"
FAM, CAT, ci = v2.FAM, v2.CAT, v2.ci
SEQ = {"tp": "TP", "pvp": "PVP", "dwi_b2": "DWI-b2", "hbp": "HBP", "ap": "AP", "adc": "ADC", "pre": "Pre", "t2": "T2", "dwi_b1": "DWI-b1"}
CLS = {"glcm": "GLCM", "gldm": "GLDM", "glrlm": "GLRLM", "glszm": "GLSZM", "ngtdm": "NGTDM", "firstorder": "First-order", "shape": "Shape"}
CLIN = {"age": "Age", "male": "Sex (male)", "bmi": "BMI", "etiology_hbv": "HBV etiology", "etiology_hcv": "HCV etiology", "cirrhosis": "Cirrhosis", "log_afp": "log AFP", "tumor_size_pvp_cm": "Tumor size (PVP)",
        "capsule": "Capsule", "margin_irregular": "Irregular margin", "mosaic": "Mosaic", "internal_arteries": "Internal arteries", "fat": "Fat content", "t2_hyper": "T2 hyperintense", "hbp_hypo": "HBP hypointense",
        "lirads5": "LI-RADS 5", "treat_io": "Treatment: IO", "treat_io_sbrt": "Treatment: IO + SBRT", "grade": "Tumor grade", "macrovascular_invasion": "Macrovascular invasion"}


def ep_cfg(key, fb=False):
    t, g = BASE / "main_pipeline_tils", BASE / "main_pipeline_tls"
    c = {"tils_ge2": ("TILs 2-3 vs 0-1", t, "ladder_post_20261021/tilsGE2_post_*", "grade_ge2_clinrad_post_20261022/tilsGE2cr_post_*", "grade_ge2_clinical_only_post_20261022"),
         "tils_ge3": ("TILs 3 vs 0-2", t, "ge3_ladder_post_20261021/tilsGE3_post_*", "grade_ge3_clinrad_post_20261022/tilsGE3cr_post_*", "grade_ge3_clinical_only_post_20261022"),
         "tls_ge2": ("TLS 2 vs 0-1", g, "ladder_post_20261021/tlsGE2_post_*", "grade_ge2_clinrad_post_20261022/tlsGE2cr_post_*", "grade_ge2_clinical_only_post_20261022")}[key]
    if fb:      # TEST ONLY: no-grade results
        old = {"tils_ge2": ("ladder_clinrad_post_20261021/tilsGE2cr_post_*", "clinical_only_v4_20261021"), "tils_ge3": ("ge3_ladder_clinrad_post_20261021/tilsGE3cr_post_*", "ge3_clinical_only_post_20261021"),
               "tls_ge2": ("ladder_clinrad_post_20261021/tlsGE2cr_post_*", "clinical_only_post_20261021")}[key]; c = (c[0], c[1], c[2], old[0], old[1])
    D = OLD / key if fb else NEW / key
    return dict(label=c[0], prefix=key, dirs={"radiomics": OLD / key, "radiomics_clinical": D, "clinical": D},
                oof={"radiomics": str(c[1] / c[2] / "oof_predictions.csv"), "radiomics_clinical": str(c[1] / c[3] / "oof_predictions.csv"), "clinical": str(c[1] / c[4] / "oof_predictions.csv")})


def pretty(f):
    if f.startswith("clin_"): return "Clinical · " + CLIN.get(f[5:], f[5:])
    if ":" in f:
        seq, rest = f.split(":", 1); cls, name = rest.split("_", 1); return f"{SEQ.get(seq, seq)} · {CLS.get(cls, cls)} · " + re.sub(r"(?<!^)(?=[A-Z])", " ", name)
    return "Clinical · " + CLIN.get(f, f)


# -------------------------------------------------------------- selection
def add_score(df, kind):
    df = df.copy(); df["score"] = (df.oof_auc_avgprob + df.lock_auc) / 2 if kind == "cls" else (df.oof_c_avg + df.lock_c) / 2; return df


def per_combo(df):
    return df.loc[df.groupby("sequence_set").score.idxmax()].sort_values("score", ascending=False).reset_index(drop=True)


def load_tables(cfg, prefix, kind, fallback):
    out = {}
    for tag, _, _ in CAT:
        f = cfg["dirs"][tag] / f"{prefix}_model_table_{tag}.csv"
        if not f.exists() and fallback: f = (OLD / cfg["prefix"] if kind == "cls" else OLD / "recurrence") / f"{prefix}_model_table_{tag}.csv"
        if not f.exists(): return None
        out[tag] = add_score(pd.read_csv(f), kind)
    return out


def cls_rows(df, n=None):
    rows = []
    for _, r in (df.head(n) if n else df).iterrows():
        deg = min(r.lock_sens, r.lock_spec) <= 0.05
        rows.append({"Sequences": r.sequence_set.replace("+", " + ") if r.sequence_set != "clinicoradiologic" else "all variables", "Model": FAM[r.family] + (" †" if deg else ""), "Mean of OOF & lock AUC": f"{r.score:.2f}",
                     "OOF AUC (95% CI)": ci(r.oof_auc_avgprob, r.oof_auc_ci_lo, r.oof_auc_ci_hi), "Lock AUC (95% CI)": ci(r.lock_auc, r.lock_auc_ci_lo, r.lock_auc_ci_hi),
                     "Lock sensitivity (95% CI)": ci(r.lock_sens, r.lock_sens_ci_lo, r.lock_sens_ci_hi), "Lock specificity (95% CI)": ci(r.lock_spec, r.lock_spec_ci_lo, r.lock_spec_ci_hi)})
    return pd.DataFrame(rows)


def rec_rows(df, n=None):
    rows = []
    for _, r in (df.head(n) if n else df).iterrows():
        rows.append({"Sequences": r.sequence_set.replace("+", " + ") if r.sequence_set != "clinical" else "all variables", "Model": FAM[r.family], "Mean of OOF & lock C-index": f"{r.score:.2f}",
                     "OOF C-index (95% CI)": ci(r.oof_c_avg, r.oof_c_ci_lo, r.oof_c_ci_hi), "Lock C-index (95% CI)": ci(r.lock_c, r.lock_c_ci_lo, r.lock_c_ci_hi)})
    return pd.DataFrame(rows)


SEL_FOOT = ("Model selection: one model per sequence combination, chosen as the family with the highest mean of OOF and lock {m}; combinations ranked by the same score. OOF = {m} of the split-averaged out-of-fold predictions (development set, "
            "10 repeated 5-fold splits; bootstrap 95% CI). Because the lock set enters the selection score, lock estimates of selected models are NOT independent validation.")


def perf_slide(prs, title, df, kind, extra=""):
    s = bd.new_slide(prs, title); n = len(df)
    widths = [1.9, 1.2, 1.2, 2.0, 2.0, 2.0, 2.0] if kind == "cls" else [2.4, 1.4, 1.6, 3.0, 3.0]
    bd.add_table(s, df, 0.45 if kind == "cls" else 0.9, 1.1, widths, font=9.5 if n <= 10 else 8.5, row_h=0.4 if n <= 10 else 0.33, align_center_from=2, header_h=0.55)
    bd.footnote(s, SEL_FOOT.format(m="AUC" if kind == "cls" else "C-index") + (" Lock sensitivity / specificity at the Youden threshold of the dev OOF predictions; † = (almost) everyone classified one way. " if kind == "cls" else " ") + extra, top=6.8)


# -------------------------------------------------------------- figures
def model_folder(tag, r):
    return f"{r.sequence_set.replace('+', '_')}__{r.family}" if tag not in ("clinical",) else f"clinicoradiologic__{r.family}"


def roc_fig(cfg, best, path):
    fig, ax = plt.subplots(1, 2, figsize=(12.5, 4.9)); handles = []
    for tag, name, col in CAT:
        r = best[tag]; o = pd.concat([pd.read_csv(f) for f in run_guard.usable_files(cfg["oof"][tag], verbose=False)]); o = o[(o.family == r.family) & o.prob.notna()]
        if tag != "clinical": o = o[o.sequence_set == r.sequence_set]
        avg = o.groupby("PatientID").agg(y=("y", "first"), prob=("prob", "mean")); fpr, tpr, _ = roc_curve(avg.y, avg.prob)
        h, = ax[0].plot(fpr, tpr, color=col, lw=2, label=f"{name} ({r.sequence_set.replace('+', ' + ') if tag != 'clinical' else 'all variables'}, {FAM[r.family].lower()})"); handles.append(h)
        lp = pd.read_csv(cfg["dirs"][tag] / "models" / tag / model_folder(tag, r) / "lock_predictions.csv"); fpr, tpr, _ = roc_curve(lp.y_lock, lp.prob); ax[1].plot(fpr, tpr, color=col, lw=2)
    for k, t in enumerate(("Development set: out-of-fold (averaged over splits)", "Lock set (one-shot; also used for model selection)")):
        ax[k].plot([0, 1], [0, 1], "--", color="gray", lw=1); ax[k].set_xlabel("1 - specificity"); ax[k].set_ylabel("Sensitivity"); ax[k].set_title(t, fontsize=11); ax[k].set_aspect("equal"); ax[k].spines[["top", "right"]].set_visible(False)
    fig.legend(handles=handles, loc="lower center", ncol=3, frameon=False, fontsize=9.5); plt.tight_layout(rect=(0, 0.07, 1, 1)); plt.savefig(path, dpi=200, bbox_inches="tight"); plt.close()


def shap_top(dirs, tag, r, n=10):
    f = dirs[tag] / "models" / tag / model_folder(tag, r) / "shap_coefficients.csv"
    return pd.read_csv(f).sort_values("mean_abs_shap", ascending=False) if f.exists() else None


def shap_slide(prs, title, best, dirs, lists):
    s = bd.new_slide(prs, title); x = 0.3
    for tag, name, _ in CAT:
        r = best[tag]; sh = shap_top(dirs, tag, r)
        label = f"{name}: {r.sequence_set.replace('+', ' + ') if tag != 'clinical' else 'all variables'} / {FAM[r.family].lower()}" + (f" (n={len(sh)} features)" if sh is not None else "")
        bd.text_box(s, x, 0.95, 4.2, 0.5, [(label,)], size=10)
        if sh is not None:
            lists.append((title.split(" - ")[0], name, label, sh))
            t = pd.DataFrame({"Feature": [pretty(f) for f in sh.feature.head(10)], "Mean |SHAP|": [f"{v:.3f}" for v in sh.mean_abs_shap.head(10)]})
            bd.add_table(s, t, x, 1.5, [3.3, 0.85], font=8.5, row_h=0.36, align_center_from=1)
        x += 4.35
    bd.footnote(s, "Mean absolute linear SHAP values in the final development-fit models (standardized features; top 10 shown; scales are model-specific). Complete feature lists follow in the appendix.", top=6.9)


def feature_list_slide(prs, endpoint, items):
    s = bd.new_slide(prs, f"{endpoint} - features in plotted models"); x = 0.3
    for _, cat, label, sh in items:
        lines = [(label,)] + [f"{i + 1}. {pretty(f)}" for i, f in enumerate(sh.feature)]
        bd.text_box(s, x, 0.95, 4.25, 6.2, lines, size=6.5 if len(sh) > 22 else 7.5); x += 4.35


# -------------------------------------------------------------- sections
def classification_section(prs, key, D, split, spec_extra, fallback, lists, summary):
    cfg = ep_cfg(key, fallback); tabs = load_tables(cfg, key, "cls", fallback)
    if tabs is None: print("skip", key, "(tables missing)"); return None
    dev_n, lock_n = int((split["set"] == "dev").sum()), int((split["set"] == "lock").sum())
    spec = [("Age", "Age, years", "cont"), ("Gender", "Gender", "cat"), ("BMI", "BMI, kg/m2", "cont"), ("Cirrhosis", "Cirrhosis", "cat"), ("Capsule (on PVP yes/no)", "Capsule (on PVP)", "cat"),
            ("Mosaic appearance (yes/no)", "Mosaic appearance", "cat"), ("Macrovascular invasion (yes/no)", "Macrovascular invasion", "cat"), ("Maximum tumor size on T1 portal venous phase (PVP)", "Maximum tumor size on T1 PVP, cm", "cont")] + spec_extra
    v2.split_slide(prs, f"{cfg['label']} - development/testing data split", v2.split_table(D, split, spec, dev_n, lock_n))
    rb, cb = per_combo(tabs["radiomics"]), per_combo(tabs["radiomics_clinical"]); clb = tabs["clinical"].sort_values("score", ascending=False).reset_index(drop=True)
    perf_slide(prs, f"{cfg['label']} - radiomics models, top 10 of 14 sequence combinations", cls_rows(rb, 10), "cls")
    perf_slide(prs, f"{cfg['label']} - radiomics + clinico-radiologic variables (incl. tumor grade), all 14 combinations", cls_rows(cb), "cls")
    best = {"radiomics": rb.iloc[0], "radiomics_clinical": cb.iloc[0], "clinical": clb.iloc[0]}
    s = bd.new_slide(prs, f"{cfg['label']} - best model of each category")
    t = cls_rows(pd.DataFrame([best[t] for t, _, _ in CAT])); t.insert(0, "Category", [n for _, n, _ in CAT])
    bd.add_table(s, t, 0.3, 1.05, [2.1, 1.25, 0.9, 1.05, 1.85, 1.85, 1.8, 1.8], font=8.5, row_h=0.42, align_center_from=3, header_h=0.55)
    p = (NEW if not fallback else NEW) / key; p.mkdir(parents=True, exist_ok=True); roc_fig(cfg, best, p / "roc_best_by_mean.png")
    pic = bd.picture(s, p / "roc_best_by_mean.png", 0.9, 3.0, height=3.95); pic.left = int((prs.slide_width - pic.width) / 2)
    bd.footnote(s, "Best per category = highest mean of OOF and lock AUC. Left: OOF predictions averaged over splits (n=%d). Right: lock set (n=%d); the lock set also entered the selection." % (dev_n, lock_n), top=7.05)
    shap_slide(prs, f"{cfg['label']} - mean absolute SHAP", best, cfg["dirs"], lists); summary[key] = (cfg["label"], best, tabs)
    return best


def recurrence_section(prs, D, split, fallback, lists, summary):
    D2 = OLD / "recurrence" if fallback else NEW / "recurrence"; cfg = dict(dirs={"radiomics": OLD / "recurrence", "radiomics_clinical": D2, "clinical": D2})
    tabs = load_tables(cfg, "recurrence", "rec", fallback)
    if tabs is None: print("skip recurrence"); return None
    dev_n, lock_n = int((split["set"] == "dev").sum()), int((split["set"] == "lock").sum())
    spec = [("Age", "Age, years", "cont"), ("Gender", "Gender", "cat"), ("Cirrhosis", "Cirrhosis", "cat"), ("Capsule (on PVP yes/no)", "Capsule (on PVP)", "cat"), ("Mosaic appearance (yes/no)", "Mosaic appearance", "cat"),
            ("Macrovascular invasion (yes/no)", "Macrovascular invasion", "cat"), ("Maximum tumor size on T1 portal venous phase (PVP)", "Maximum tumor size on T1 PVP, cm", "cont"), ("Serum AFP", "Serum AFP, ng/mL", "cont"),
            ("Recurrence", "Recurrence", "cat"), ("Follow-up time (days)", "Follow-up / time to recurrence, days", "cont")]
    v2.split_slide(prs, "Recurrence - development/testing data split", v2.split_table(D, split, spec, dev_n, lock_n), note="Follow-up time = time to recurrence or last follow-up. One patient without follow-up excluded.")
    rb, cb = per_combo(tabs["radiomics"]), per_combo(tabs["radiomics_clinical"]); clb = tabs["clinical"].sort_values("score", ascending=False).reset_index(drop=True)
    perf_slide(prs, "Recurrence - radiomics models, top 10 of 14 sequence combinations", rec_rows(rb, 10), "rec", "The lock set was opened earlier in this project.")
    perf_slide(prs, "Recurrence - radiomics + clinico-radiologic variables (incl. tumor grade), all 14 combinations", rec_rows(cb), "rec", "The lock set was opened earlier in this project.")
    best = {"radiomics": rb.iloc[0], "radiomics_clinical": cb.iloc[0], "clinical": clb.iloc[0]}
    R = BASE / "main_pipeline_recurrence"; km_dir = NEW / "recurrence" / "km"
    spec_json = {"best": [dict(tag=t, sequences=best[t].sequence_set, family=best[t].family) for t, _, _ in CAT], "model_dirs": {k: str(v) for k, v in cfg["dirs"].items()},
                 "oof": {"radiomics": str(R / "ladder14_radiomics_post_20261021/recA_post_none_*/oof_risk.csv"), "radiomics_clinical": str(R / "grade_ladder14_clinrad_post_20261022/recB_post_forced_*/oof_risk.csv"), "clinical": str(R / "clinical/clinsurv_pre_grade_*/oof_risk.csv")}, "out": str(km_dir)}
    if fallback: spec_json["oof"]["radiomics_clinical"] = str(R / "ladder14_radiomics_clinical_post_20261021/recB_post_forced_*/oof_risk.csv"); spec_json["oof"]["clinical"] = str(R / "clinical/clinsurv_pre_2*/oof_risk.csv")
    sj = NEW / "recurrence" / "km_spec.json"; sj.parent.mkdir(parents=True, exist_ok=True); sj.write_text(json.dumps(spec_json, indent=1))
    subprocess.run(["/usr/bin/python3", str(Path(__file__).resolve().parent / "make_km_figures.py"), "--spec", str(sj)], check=True, capture_output=True)
    km = pd.read_csv(km_dir / "km_summary.csv")
    for setname, title in (("lock", "lock set (also used for selection)"), ("dev", "development set, out-of-fold risk")):
        s = bd.new_slide(prs, f"Recurrence - best model of each category: {title}"); x = 0.3
        for tag, name, _ in CAT:
            r = best[tag]; k = km[(km.tag == tag) & (km.set == setname)].iloc[0]
            row = pd.DataFrame([{"Model": f"{name}: {r.sequence_set.replace('+', ' + ') if tag != 'clinical' else 'all variables'}, {FAM[r.family].lower()}",
                                 "C-index (95% CI)": ci(r.oof_c_avg, r.oof_c_ci_lo, r.oof_c_ci_hi) if setname == "dev" else ci(r.lock_c, r.lock_c_ci_lo, r.lock_c_ci_hi), "Log-rank p": "<0.001" if k.logrank_p < 0.001 else f"{k.logrank_p:.3f}"}])
            bd.add_table(s, row, x, 1.05, [2.0, 1.2, 0.8], font=8, row_h=0.62, align_center_from=1, header_h=0.3)
            bd.picture(s, km_dir / f"km_{tag}_{setname}.png", x, 2.05, width=4.0); bd.picture(s, km_dir / f"tdauc_{tag}_{setname}.png", x + 0.25, 5.1, width=3.5); x += 4.35
        bd.footnote(s, "Best per category = highest mean of OOF and lock C-index. High / low risk = above / at-or-below the median risk of the set shown. Time-dependent AUC(t): censoring-weighted. Lock set was opened earlier and enters selection: exploratory.", top=7.1)
    shap_slide(prs, "Recurrence - mean absolute SHAP", best, cfg["dirs"], lists); summary["rec"] = ("Recurrence", best, tabs)
    return best


def grade_slide(prs, summary):
    rows = []
    for key, (label, best, tabs) in summary.items():
        for tag, name in (("clinical", "Clinico-radiologic"), ("radiomics_clinical", "Radiomics + clinico-radiologic")):
            kind = "rec" if key == "rec" else "cls"; old_dir = OLD / ("recurrence" if key == "rec" else key)
            f = old_dir / (f"recurrence_model_table_{tag}.csv" if key == "rec" else f"{key}_model_table_{tag}.csv")
            if not f.exists(): continue
            o = add_score(pd.read_csv(f), kind); o = o.loc[o.score.idxmax()]; n = best[tag]
            g = lambda r: (f"{r.oof_c_avg:.2f}", f"{r.lock_c:.2f}") if kind == "rec" else (f"{r.oof_auc_avgprob:.2f}", f"{r.lock_auc:.2f}")
            sq = lambda r: r.sequence_set.replace("+", " + ") if r.sequence_set not in ("clinical", "clinicoradiologic") else "all"
            rows.append({"Endpoint": label, "Category": name, "No grade: model": f"{sq(o)} / {FAM[o.family].lower()}", "OOF": g(o)[0], "Lock": g(o)[1], "With grade: model": f"{sq(n)} / {FAM[n.family].lower()}", "OOF ": g(n)[0], "Lock ": g(n)[1]})
    s = bd.new_slide(prs, "Tumor grade sensitivity analysis")
    bd.add_table(s, pd.DataFrame(rows), 0.4, 1.1, [1.8, 2.3, 2.0, 0.7, 0.7, 2.0, 0.7, 0.7], font=9, row_h=0.42, align_center_from=3, header_h=0.5)
    bd.footnote(s, "Best model per category (highest mean of OOF and lock AUC / C-index) without vs with tumor grade (mastersheet column BM, ordinal G1-G4; 'G2/G3' = 2.5; imputed / scaled within folds). OOF / Lock = AUC (C-index for recurrence). Lock set enters selection: exploratory.", top=6.6)


def selection_slide(prs):
    s = bd.new_slide(prs, "Model selection")
    bd.text_box(s, 0.5, 1.0, 12.3, 5.8, [("Selection rule",), "Best model = highest mean of development OOF performance and lock performance: (OOF AUC + lock AUC) / 2 for classification; (OOF C-index + lock C-index) / 2 for recurrence.",
        "OOF AUC = AUC of the out-of-fold predictions averaged over 10 repeated 5-fold splits (the value in the tables). The rule picks the family for each sequence combination, ranks the combinations, and picks the best model of each category.",
        ("Interpretation",), "The lock set (n=30) is part of the selection, so lock estimates of the selected models are selected estimates and should not be described as independent validation; the OOF estimate is also from the data used for selection.",
        "With 14 combinations x 3 families per category and a small lock set (7 TILs<2 and 20 TILs<3 negatives), the selected numbers are optimistic. Lock-set confidence intervals are wide.",
        "The recurrence lock set had already been opened earlier in this project."], size=14, bullet=True)


def conclusions(prs, summary):
    lines = []
    for key, (label, best, tabs) in summary.items():
        r, c, k = best["radiomics"], best["radiomics_clinical"], best["clinical"]; m = lambda x, rec=(key == "rec"): (x.oof_c_avg, x.lock_c) if rec else (x.oof_auc_avgprob, x.lock_auc)
        nm = lambda x: x.sequence_set.replace("+", " + ") if x.sequence_set not in ("clinical", "clinicoradiologic") else "all variables"
        lines += [(label,), f"Radiomics {nm(r)} ({FAM[r.family].lower()}): OOF {m(r)[0]:.2f}, lock {m(r)[1]:.2f}; radiomics + clinico-radiologic {nm(c)}: OOF {m(c)[0]:.2f}, lock {m(c)[1]:.2f}; clinico-radiologic alone: OOF {m(k)[0]:.2f}, lock {m(k)[1]:.2f}."]
    lines += [("Caveats",), "Models selected on the mean of OOF and lock performance: lock numbers are selected estimates, not independent validation. Category differences are small relative to the 95% CIs; lock sets are small."]
    s = bd.new_slide(prs, "Conclusions"); bd.text_box(s, 0.5, 1.0, 12.3, 5.9, lines, size=12, bullet=True)
    s = bd.new_slide(prs, "Next steps...")
    bd.text_box(s, 0.5, 1.0, 12.3, 5.0, ["Extend the cohort (extended analysis, N=220) for external-style validation of frozen models selected without the lock set.", "Target IHC-derived endpoints.", "Target assessment of response to ICIs in SA2."], size=14, bullet=True)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--allow-fallback", action="store_true", help="TEST ONLY: use the no-grade results when the grade results do not exist yet"); a = ap.parse_args()
    prs = Presentation(bd.TEMPLATE); ids = prs.slides._sldIdLst
    for sid in list(ids)[1:]: prs.part.drop_rel(sid.rId); ids.remove(sid)
    bd.text_box(prs.slides[0], 0.26, 5.6, 10, 0.5, [f"Results update, {date.today():%m/%d/%y}"], size=16)
    D, r = bt.build_frame(); tils = pd.read_csv(BASE / "main_pipeline/final/cohort_v3_141/split/split.csv").set_index("PatientID"); rec = pd.read_csv(BASE / "main_pipeline_recurrence/split/recurrence_split.csv").set_index("PatientID")
    s = bd.new_slide(prs, "Patients")
    pt = pd.DataFrame([{"Specific aim": "SA1", "Cohort size": "N = 111", "Description": "Retrospective cohort of resected, untreated HCC patients"}, {"Specific aim": "SA2", "Cohort size": "N = 30", "Description": "HCC patients who received neoadjuvant immunotherapy (IO, n=16; IO + SBRT, n=14) prior to resection"}])
    bd.add_table(s, pt, 0.5, 1.0, [2.0, 2.0, 7.6], font=12, row_h=0.5, align_center_from=1)
    bd.text_box(s, 0.5, 2.5, 11.5, 3.5, [("Inclusion criteria",), "Adult patients with chronic liver disease of any etiology and resected HCC (from last 14 years)", "Contrast-enhanced MRI (using gadoxetate) within 3 months of resection", "HCC lesions ≥ 2 cm", "Treatment with neoadjuvant IO (SA2 only)",
                                         ("Exclusion criteria",), "Poor MR image quality; previous locoregional or systemic therapy; unavailable tissue; previous liver transplant"], size=13, bullet=True)
    s = bd.new_slide(prs, "Data split and training")
    bd.text_box(s, 0.5, 1.1, 12.3, 5.6, [("Split",), "TILs / TLS: 111 / 30 patients in development / lock sets (same split for both). Recurrence: 110 / 30; one patient without follow-up excluded.", "Lock-set use: see the Model selection slide (the lock set enters the selection score).",
                                         ("Validation (development set)",), "10 x repeated stratified 5-fold cross-validation (inner 3-fold for tuning); every preprocessing, selection and tuning step is refit inside each training fold.", "Performance = AUC (classification) or Harrell C-index (recurrence) of pooled out-of-fold predictions, with 95% CIs.",
                                         ("Endpoints",), "TILs 2-3 vs 0-1;  TILs 3 vs 0-2;  TLS 2 vs 0-1;  recurrence (time to recurrence; 2-year recurrence for ROC curves)."], size=14, bullet=True)
    s = bd.new_slide(prs, "Radiomics model development")
    bd.text_box(s, 0.5, 1.0, 7.9, 5.9, [("Per training fold",), "963 original PyRadiomics features (9 sequences x 107), per-sequence bin widths, N4-corrected images; SA2 imaging is post-treatment.", "Inter-reader filter: ICC(3,1) ≥ 0.85 on 23 patients segmented by two readers (estimated once on development data).",
                                         "QC: ≥ 80% observed, CV > 0.10; median imputation.", "Within-sequence complete-linkage clustering, |Spearman ρ| ≥ 0.85; one outcome-ranked representative per cluster.", "Ridge (top 10 / 15 / 20 clusters), lasso and elastic net (10-20 non-zero coefficients) logistic / Cox models; regularization tuned by inner 3-fold CV.",
                                         ("Clinico-radiologic variables",), "Age, sex, BMI, etiology, cirrhosis, log AFP, tumor size, qualitative MRI (capsule, margin, mosaic, internal arteries, fat, T2, HBP, LI-RADS), treatment (untreated / IO / IO + SBRT) and tumor grade (pathology, ordinal G1-G4; mastersheet column BM). Binary variables with minority < 10% dropped.",
                                         "Combined models: clinical variables added to the selected radiomics features (ridge) or to the candidate pool (lasso / elastic net)."], size=11.5, bullet=True)
    combos = pd.DataFrame({"Sequence combinations (14)": ["TP", "PVP", "DWI (b2)", "HBP", "TP + PVP", "TP + DWI", "TP + HBP", "PVP + DWI", "PVP + HBP", "DWI + HBP", "TP + PVP + DWI", "TP + PVP + HBP", "TP + DWI + HBP", "PVP + DWI + HBP"]})
    bd.add_table(s, combos, 9.0, 1.1, [3.5], font=10, row_h=0.33, align_center_from=9)
    v2.cohort_slide(prs, D); selection_slide(prs)
    lists, summary = [], {}
    ex = [("TILs_resection", "TILs score", "cat"), ("TLS_resection", "TLS score", "cat")]
    classification_section(prs, "tils_ge2", D, tils, ex, a.allow_fallback, lists, summary); classification_section(prs, "tils_ge3", D, tils, ex, a.allow_fallback, lists, summary)
    classification_section(prs, "tls_ge2", D, tils, ex[::-1], a.allow_fallback, lists, summary); recurrence_section(prs, D, rec, a.allow_fallback, lists, summary)
    grade_slide(prs, summary)
    seen = {}
    for ep, cat, label, sh in lists: seen.setdefault(ep, []).append((ep, cat, label, sh))
    for ep, items in seen.items(): feature_list_slide(prs, ep, items)
    pd.concat([sh.assign(endpoint=ep, category=cat, model=label) for ep, cat, label, sh in lists]).to_csv(NEW / "plotted_model_feature_lists.csv", index=False) if lists else None
    conclusions(prs, summary)
    dest = NEW / f"updates-{date.today():%m-%d-%y}_claude_v3{'_TEST' if a.allow_fallback else ''}.pptx"; NEW.mkdir(parents=True, exist_ok=True); prs.save(dest); print("saved", dest, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
